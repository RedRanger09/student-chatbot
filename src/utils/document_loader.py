"""
Document loading helpers for institutional and session knowledge bases.

Supports PDF, DOCX, TXT, and Markdown. Unsupported or corrupt files are skipped
with a warning rather than crashing the pipeline.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from src.utils.logger import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset({".pdf", ".docx", ".txt", ".md"})


@dataclass(frozen=True)
class LoadedDocument:
    """A successfully loaded source document."""

    text: str
    document_name: str
    source_path: str


def is_supported(path: Path) -> bool:
    """Return True if ``path`` has a supported document extension."""
    return path.suffix.lower() in SUPPORTED_EXTENSIONS


def load_document(path: Path, *, document_name: str | None = None) -> LoadedDocument | None:
    """
    Load text from a single PDF, DOCX, TXT, or Markdown file.

    Args:
        path: Path to the document on disk.
        document_name: Optional display name (e.g. relative path under KB root).

    Returns:
        A ``LoadedDocument`` on success, or ``None`` if the file is
        unsupported, empty, or unreadable.
    """
    if not path.is_file():
        logger.warning("Skipping non-file path: %s", path)
        return None

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        logger.warning("Ignoring unsupported file: %s", path.name)
        return None

    try:
        if suffix == ".pdf":
            text = _load_pdf(path)
        elif suffix == ".docx":
            text = _load_docx(path)
        else:
            text = _load_txt(path)
    except Exception as exc:  # noqa: BLE001 — isolate corrupt files
        logger.error("Failed to load %s: %s", path.name, exc)
        return None

    text = (text or "").strip()
    if not text:
        logger.warning("Empty document skipped: %s", path.name)
        return None

    logger.info("Loaded document: %s (%s chars)", path.name, len(text))
    return LoadedDocument(
        text=text,
        document_name=document_name or path.name,
        source_path=str(path.resolve()),
    )


def load_documents_from_directory(
    directory: Path,
    *,
    recursive: bool = True,
) -> list[LoadedDocument]:
    """
    Load all supported documents from ``directory``.

    When ``recursive`` is True (default), subdirectories are scanned and
  document names use paths relative to ``directory`` (e.g. ``hostel/hostel_faq.md``).

    Missing or empty directories yield an empty list without raising.
    """
    if not directory.exists():
        logger.warning("Documents folder missing: %s", directory)
        return []

    if not directory.is_dir():
        logger.error("Documents path is not a directory: %s", directory)
        return []

    documents: list[LoadedDocument] = []
    paths = (
        sorted(p for p in directory.rglob("*") if p.is_file())
        if recursive
        else sorted(p for p in directory.iterdir() if p.is_file())
    )
    for path in paths:
        if not is_supported(path):
            if path.name not in {".gitkeep", "Thumbs.db", ".DS_Store"}:
                logger.warning("Ignoring unsupported file: %s", path)
            continue
        rel_name = path.relative_to(directory).as_posix()
        loaded = load_document(path, document_name=rel_name)
        if loaded is not None:
            documents.append(loaded)

    logger.info("Documents loaded from %s: %s", directory, len(documents))
    return documents


def _load_pdf(path: Path) -> str:
    """Extract text from a PDF using pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page_num, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            logger.warning("Corrupt/unreadable PDF page %s in %s: %s", page_num, path.name, exc)
            continue
        if page_text.strip():
            parts.append(page_text)
    return "\n\n".join(parts)


def _load_docx(path: Path) -> str:
    """Extract text from a DOCX using python-docx."""
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n\n".join(parts)


def _load_txt(path: Path) -> str:
    """Read a plain-text file with UTF-8 (fallback to latin-1)."""
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        logger.warning("UTF-8 decode failed for %s; retrying latin-1", path.name)
        return path.read_text(encoding="latin-1", errors="replace")
